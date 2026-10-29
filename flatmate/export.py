"""Excel review view of one flat. SQLite stays the authoritative store: this file is
written fresh each time and never read back, so editing it cannot change any record."""

import json

from openpyxl import Workbook
from openpyxl.styles import Font

from flatmate import balances, config, ledger
from flatmate.db import log, today


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(list(row))
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(
            60, max(10, *(len(str(c.value or "")) + 2 for c in col)))


def export_flat(conn, flat_id):
    ledger.get_flat(conn, flat_id)
    rs = lambda paisa: paisa / 100  # rupees as a number, so Excel can sum it
    wb = Workbook()
    wb.remove(wb.active)

    net = balances.net_balances(conn, flat_id)
    _sheet(wb, "Balances", ["Member", f"Balance ({config.CURRENCY})", "Meaning"],
           [(m, rs(v), "is owed" if v > 0 else "owes" if v < 0 else "settled") for m, v in net.items()])

    _sheet(wb, "Expenses (all versions)",
           ["Expense", "Version", "Status", "Date", "Description", "Category", f"Amount ({config.CURRENCY})",
            "Paid by", "Split", "Logged by", "Approved by", "Reason", "Source message"],
           [(e["expense_id"], e["version"], e["status"], e["expense_date"], e["description"], e["category"],
             rs(e["amount"]), e["payer_id"],
             ", ".join(f"{s['member_id']}={rs(s['share'])}" for s in conn.execute(
                 "SELECT member_id, share FROM expense_shares WHERE row_id = ?", (e["row_id"],))),
             e["created_by"], e["approved_by"], e["reason"], e["source_msg_id"])
            for e in conn.execute("SELECT * FROM expenses WHERE flat_id = ? ORDER BY expense_id, version",
                                  (flat_id,))])

    _sheet(wb, "Evidence", ["Expense", "Kind", "Message", "Channel", "From", "Text"],
           conn.execute("SELECT ev.expense_id, ev.kind, ev.msg_id, m.channel, m.sender_id, m.text "
                        "FROM expense_evidence ev JOIN messages m ON m.msg_id = ev.msg_id "
                        "WHERE m.flat_id = ? ORDER BY ev.expense_id", (flat_id,)).fetchall())

    _sheet(wb, "Preferences", ["Id", "Member", "Rule", "Category", "Valid from", "Status", "Supersedes",
                               "Proposed by", "Approved by", "Note"],
           conn.execute("SELECT pref_id, member_id, key, value, valid_from, status, supersedes_id, created_by, "
                        "approved_by, note FROM preferences WHERE flat_id = ? ORDER BY pref_id", (flat_id,)))

    _sheet(wb, "Payments", ["Id", "From", "To", f"Amount ({config.CURRENCY})", "Status", "Confirmed by"],
           [(s["settlement_id"], s["from_id"], s["to_id"], rs(s["amount"]), s["status"], s["confirmed_by"])
            for s in conn.execute("SELECT * FROM settlements WHERE flat_id = ?", (flat_id,))])

    _sheet(wb, "Messages", ["Message", "Date", "Channel", "From", "Text", "Outcome"],
           conn.execute("SELECT msg_id, business_date, channel, sender_id, text, outcome FROM messages "
                        "WHERE flat_id = ? ORDER BY received_at", (flat_id,)))

    _sheet(wb, "Outbox", ["Id", "Kind", "To", "Subject", "Approval", "Status", "Attempts", "Error", "Sent at"],
           conn.execute("SELECT outbox_id, kind, recipient_id, subject, approval, status, attempts, last_error, "
                        "sent_at FROM outbox WHERE flat_id = ? ORDER BY outbox_id", (flat_id,)))

    _sheet(wb, "Action log", ["Time", "Cycle", "Actor", "Event", "Tool", "Status", "Error", "Input", "Output"],
           conn.execute("SELECT ts, cycle_id, actor, event, tool, status, error, input, output FROM action_log "
                        "WHERE flat_id = ? ORDER BY log_id", (flat_id,)))

    config.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    path = config.EXPORT_DIR / f"{flat_id}_{today(conn).isoformat()}.xlsx"
    wb.save(path)

    # The same action log as plain JSON lines, one event per line, for reading without Excel.
    log_path = path.with_name(f"{flat_id}_{today(conn).isoformat()}_action_log.jsonl")
    with open(log_path, "w", encoding="utf-8") as f:
        for r in conn.execute("SELECT * FROM action_log WHERE flat_id = ? ORDER BY log_id", (flat_id,)):
            row = dict(r)
            for key in ("input", "output"):
                row[key] = json.loads(row[key]) if row[key] else None
            f.write(json.dumps(row, ensure_ascii=False) + "\n")

    log(conn, "tool_call", flat_id=flat_id, tool="export_excel", actor="user",
        output={"file": str(path), "log": str(log_path)})
    return path
